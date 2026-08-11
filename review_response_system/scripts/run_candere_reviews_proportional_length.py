"""
Re-run of run_candere_reviews.py with two changes:
  - LLM: Claude Sonnet 5 (via ANTHROPIC_API_KEY), instead of Haiku 4.5.
  - Response length is no longer the fixed Default (80-word) cap. It's
    computed per review:

      * Negative reviews (1-2 stars): always 50 words, flat — not
        proportional to review length.
      * Everything else (3-5 stars): proportional to the review's character
        length, using a ~4-characters-per-word estimate (i.e. word target =
        review_char_length / 4), capped at 80 words if that would exceed it.
      * No review text: falls back to the existing rating-only default (40
        words), since there's no length to be proportional to.

Tone stays locked by rating, same mapping as the original run:
    1-2 stars -> Empathetic
    3 stars   -> Conversational
    4-5 stars -> Friendly

Usage:
    export ANTHROPIC_API_KEY=sk-ant-...
    python scripts/run_candere_reviews_proportional_length.py <input.xlsx> <output.xlsx>
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from review_response_system.core import (
    AIResponseLengthType,
    BusinessInfo,
    Keywords,
    ReviewDetails,
    ReviewResponseRequest,
    SignatureSupportDetails,
    Tone,
    generate_review_response,
)

RATING_TONE_MAP = {
    1: Tone.EMPATHETIC,
    2: Tone.EMPATHETIC,
    3: Tone.CONVERSATIONAL,
    4: Tone.FRIENDLY,
    5: Tone.FRIENDLY,
}

NEGATIVE_RATING_WORD_CAP = 50
DEFAULT_WORD_CAP = 80  # ceiling for the proportional formula
NO_TEXT_WORD_CAP = 40  # fallback when there's no review text to be proportional to
CHARS_PER_WORD = 4

MODEL = "claude-sonnet-5"

BUSINESS_INFO = BusinessInfo(business_name="Candere")
SUPPORT_DETAILS = SignatureSupportDetails(
    signature="Regards, Candere - A Kalyan Company",
    contact_number="2261066262",
    support_email="support@candere.com",
)


def compute_max_words(review_text: str | None, rating: int) -> int:
    if rating <= 2:
        return NEGATIVE_RATING_WORD_CAP
    if not review_text or not review_text.strip():
        return NO_TEXT_WORD_CAP
    proportional = max(1, round(len(review_text.strip()) / CHARS_PER_WORD))
    return min(proportional, DEFAULT_WORD_CAP)


def main() -> None:
    if len(sys.argv) != 3:
        print("Usage: python run_candere_reviews_proportional_length.py <input.xlsx> <output.xlsx>", file=sys.stderr)
        sys.exit(1)

    input_path, output_path = sys.argv[1], sys.argv[2]

    wb = openpyxl.load_workbook(input_path)
    ws = wb.active
    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {name: idx for idx, name in enumerate(header)}

    ws.cell(row=1, column=len(header) + 1, value="LockedTone")
    ws.cell(row=1, column=len(header) + 2, value="MaxWords")
    ws.cell(row=1, column=len(header) + 3, value="GeneratedResponse")

    rows = list(ws.iter_rows(min_row=2))
    total = len(rows)
    for i, row in enumerate(rows, start=1):
        rating = row[col["rating"]].value
        review_text = row[col["ReviewContent"]].value or None
        tone = RATING_TONE_MAP[rating]
        max_words = compute_max_words(review_text, rating)

        review_details = ReviewDetails(
            reviewer_name="Customer",
            rating=rating,
            review=review_text,
            ai_response_length_type=AIResponseLengthType.DEFAULT,
            locked_tone=tone,
        )
        request = ReviewResponseRequest(
            business_info=BUSINESS_INFO,
            review_details=review_details,
            keywords=Keywords(),
            support_details=SUPPORT_DETAILS,
        )

        try:
            response_text = generate_review_response(
                request, provider="anthropic", model=MODEL, max_words=max_words
            )
        except Exception as exc:
            response_text = f"ERROR: {exc}"

        ws.cell(row=row[0].row, column=len(header) + 1, value=tone.value)
        ws.cell(row=row[0].row, column=len(header) + 2, value=max_words)
        ws.cell(row=row[0].row, column=len(header) + 3, value=response_text)

        print(f"[{i}/{total}] rating={rating} tone={tone.value} max_words={max_words} -> {response_text[:60]!r}...")

    wb.save(output_path)
    print(f"Saved -> {output_path}")


if __name__ == "__main__":
    main()
