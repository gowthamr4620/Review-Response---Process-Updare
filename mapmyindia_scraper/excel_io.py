"""
Excel/CSV input and report output.

The client's sheet is never in a fixed shape - the URL column is called
"Link" one week and "Mappls URL" the next, headers sit on row 3 under a title
banner, and pincodes arrive as floats. Rather than demanding a template, this
module sniffs the header row and maps columns by alias, prints what it
decided, and lets you override any of it with --column-map. Guessing silently
is what turns a verification report into a fiction, so the mapping is always
echoed back to the operator.
"""

from __future__ import annotations

import csv
import json
import os
import re
from dataclasses import dataclass, field
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from .models import RowComparison, Verdict
from .normalize import normalize_phone, normalize_pincode, parse_coordinate

# Canonical field -> header aliases (compared after squashing to a-z0-9).
COLUMN_ALIASES: Dict[str, Tuple[str, ...]] = {
    "url": ("url", "link", "listingurl", "listinglink", "mapplsurl", "mapplslink",
            "mapmyindiaurl", "mapmyindialink", "maplink", "locationurl", "pageurl", "profileurl"),
    "business_name": ("businessname", "name", "storename", "branchname", "outletname",
                      "locationname", "store", "branch", "outlet", "listingname"),
    "address": ("address", "fulladdress", "storeaddress", "branchaddress",
                "locationaddress", "addressline", "completeaddress"),
    "pincode": ("pincode", "pin", "pincodezip", "zip", "zipcode", "postalcode", "postcode"),
    "latitude": ("latitude", "lat"),
    "longitude": ("longitude", "long", "lng", "lon"),
    "phone": ("phone", "phonenumber", "mobile", "mobilenumber", "contact",
              "contactnumber", "telephone", "landline", "storephone"),
    "coordinates": ("coordinates", "latlong", "latlng", "latitudelongitude", "geo", "geocode"),
}

EXPECTED_FIELDS = ("business_name", "address", "pincode", "latitude", "longitude", "phone")

_HEADER_SQUASH_RE = re.compile(r"[^a-z0-9]")


def _squash_header(value: Any) -> str:
    return _HEADER_SQUASH_RE.sub("", str(value or "").lower())


@dataclass
class InputRow:
    row_number: int
    url: Optional[str]
    expected: Dict[str, Any] = field(default_factory=dict)
    raw: Dict[str, Any] = field(default_factory=dict)


def _read_table(path: str, sheet: Optional[str] = None) -> List[List[Any]]:
    ext = os.path.splitext(path)[1].lower()
    if ext in (".csv", ".txt", ".tsv"):
        delimiter = "\t" if ext == ".tsv" else ","
        with open(path, "r", encoding="utf-8-sig", newline="") as f:
            return [row for row in csv.reader(f, delimiter=delimiter)]
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook[sheet] if sheet else workbook[workbook.sheetnames[0]]
    return [list(row) for row in worksheet.iter_rows(values_only=True)]


def detect_header(rows: Sequence[Sequence[Any]], scan_depth: int = 10) -> Tuple[int, Dict[str, int]]:
    """
    Find the header row and map canonical field -> column index.

    Scores each of the first `scan_depth` rows by how many cells look like
    known headers; the best-scoring row wins. This survives title banners and
    blank leading rows, which real client sheets are full of.
    """
    best_index, best_map, best_score = 0, {}, -1
    for index, row in enumerate(rows[:scan_depth]):
        mapping: Dict[str, int] = {}
        for col, cell in enumerate(row):
            squashed = _squash_header(cell)
            if not squashed:
                continue
            for field_name, aliases in COLUMN_ALIASES.items():
                if field_name in mapping:
                    continue
                if squashed in aliases or any(squashed.startswith(a) and len(squashed) - len(a) <= 4 for a in aliases):
                    mapping[field_name] = col
                    break
        score = len(mapping) + (2 if "url" in mapping else 0)
        if score > best_score:
            best_index, best_map, best_score = index, mapping, score
    return best_index, best_map


def _split_coordinates(value: Any) -> Tuple[Optional[float], Optional[float]]:
    if value in (None, ""):
        return None, None
    parts = re.split(r"[,;\s]+", str(value).strip())
    if len(parts) >= 2:
        return parse_coordinate(parts[0]), parse_coordinate(parts[1])
    return None, None


def read_input_rows(
    path: str,
    sheet: Optional[str] = None,
    column_map: Optional[Dict[str, int]] = None,
) -> Tuple[List[InputRow], Dict[str, int], int]:
    """Return (rows, resolved column map, header row index)."""
    table = _read_table(path, sheet)
    if not table:
        return [], {}, 0
    header_index, detected = detect_header(table)
    mapping = {**detected, **(column_map or {})}

    rows: List[InputRow] = []
    for offset, raw_row in enumerate(table[header_index + 1 :], start=header_index + 2):
        if not any(cell not in (None, "") for cell in raw_row):
            continue

        def cell(field_name: str) -> Any:
            index = mapping.get(field_name)
            if index is None or index >= len(raw_row):
                return None
            value = raw_row[index]
            return value.strip() if isinstance(value, str) else value

        latitude, longitude = cell("latitude"), cell("longitude")
        if latitude in (None, "") and "coordinates" in mapping:
            latitude, longitude = _split_coordinates(cell("coordinates"))

        rows.append(
            InputRow(
                row_number=offset,
                url=str(cell("url")).strip() if cell("url") else None,
                expected={
                    "business_name": cell("business_name"),
                    "address": cell("address"),
                    "pincode": normalize_pincode(cell("pincode")),
                    "latitude": parse_coordinate(latitude),
                    "longitude": parse_coordinate(longitude),
                    "phone": normalize_phone(cell("phone")),
                },
                raw={
                    "header_row": header_index + 1,
                    "raw_pincode": cell("pincode"),
                    "raw_phone": cell("phone"),
                },
            )
        )
    return rows, mapping, header_index


def describe_mapping(mapping: Dict[str, int], header_row: int) -> str:
    if not mapping:
        return "no columns recognised"
    def col_letter(i: int) -> str:
        letters = ""
        i += 1
        while i:
            i, rem = divmod(i - 1, 26)
            letters = chr(65 + rem) + letters
        return letters
    parts = [f"{name} -> col {col_letter(index)}" for name, index in sorted(mapping.items(), key=lambda kv: kv[1])]
    return f"header row {header_row + 1}; " + ", ".join(parts)


# ------------------------------------------------------------------ output

_FILLS = {
    "MATCH": "C6EFCE",
    "NEAR MATCH": "FFEB9C",
    "REVIEW": "FFEB9C",
    "INCOMPLETE": "FFEB9C",
    "MISMATCH": "FFC7CE",
    "FETCH ERROR": "FFC7CE",
    "MISSING ON MAPPLS": "E7E6E6",
    "MISSING IN SHEET": "E7E6E6",
    "NOT COMPARED": "E7E6E6",
}


def _style_workbook_sheet(worksheet, header_len: int, verdict_columns: Iterable[int]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="44546A")
    for col in range(1, header_len + 1):
        cell = worksheet.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    verdict_columns = set(verdict_columns)
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            if cell.column in verdict_columns and cell.value in _FILLS:
                cell.fill = PatternFill("solid", fgColor=_FILLS[cell.value])

    for col in range(1, header_len + 1):
        letter = get_column_letter(col)
        longest = max(
            (len(str(worksheet.cell(row=r, column=col).value or "")) for r in range(1, worksheet.max_row + 1)),
            default=10,
        )
        worksheet.column_dimensions[letter].width = min(48, max(12, longest + 2))
    worksheet.freeze_panes = "C2"
    if worksheet.max_row > 1:
        worksheet.auto_filter.ref = f"A1:{get_column_letter(header_len)}{worksheet.max_row}"


def _verification_headers() -> List[str]:
    headers = ["Row", "Listing URL", "eLoc", "Row Verdict"]
    labels = {
        "business_name": "Business Name",
        "address": "Address",
        "pincode": "PIN Code",
        "latitude": "Latitude",
        "longitude": "Longitude",
        "phone": "Phone",
    }
    for name in EXPECTED_FIELDS:
        headers += [f"{labels[name]} (Sheet)", f"{labels[name]} (Mappls)", f"{labels[name]} Verdict", f"{labels[name]} Score"]
    headers += ["Field Sources", "Fetched At", "Error"]
    return headers


def _verification_row(comparison: RowComparison) -> List[Any]:
    listing = comparison.listing
    row: List[Any] = [comparison.row_number, comparison.source_url, listing.eloc, comparison.row_verdict]
    for name in EXPECTED_FIELDS:
        c = comparison.comparisons.get(name)
        if c is None:
            row += [None, None, Verdict.NOT_COMPARED.value, None]
        else:
            row += [
                c.expected,
                c.scraped,
                c.verdict.value,
                round(c.score, 3) if isinstance(c.score, float) else c.score,
            ]
    row += [
        "; ".join(f"{k}={v}" for k, v in sorted(listing.field_sources.items())),
        listing.fetched_at,
        listing.error,
    ]
    return row


def write_verification_report(comparisons: Sequence[RowComparison], path: str, run_meta: Optional[Dict[str, Any]] = None) -> str:
    from openpyxl import Workbook

    workbook = Workbook()
    headers = _verification_headers()
    verdict_columns = [i + 1 for i, h in enumerate(headers) if h.endswith("Verdict")]

    sheet = workbook.active
    sheet.title = "Verification"
    sheet.append(headers)
    for comparison in comparisons:
        sheet.append(_verification_row(comparison))
    _style_workbook_sheet(sheet, len(headers), verdict_columns)

    flagged = [c for c in comparisons if c.row_verdict in ("MISMATCH", "REVIEW", "FETCH ERROR")]
    issues = workbook.create_sheet("Needs Attention")
    issues.append(headers)
    for comparison in flagged:
        issues.append(_verification_row(comparison))
    _style_workbook_sheet(issues, len(headers), verdict_columns)

    summary = workbook.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    summary.append(["Listings checked", len(comparisons)])
    summary.append(["Rows fully matching", sum(1 for c in comparisons if c.row_verdict == "MATCH")])
    summary.append(["Rows with a mismatch", sum(1 for c in comparisons if c.row_verdict == "MISMATCH")])
    summary.append(["Rows needing review", sum(1 for c in comparisons if c.row_verdict == "REVIEW")])
    summary.append(["Rows incomplete", sum(1 for c in comparisons if c.row_verdict == "INCOMPLETE")])
    summary.append(["Rows that failed to fetch", sum(1 for c in comparisons if c.row_verdict == "FETCH ERROR")])
    summary.append([])
    summary.append(["Field", *[v.value for v in Verdict]])
    for name in EXPECTED_FIELDS:
        counts = {v: 0 for v in Verdict}
        for comparison in comparisons:
            c = comparison.comparisons.get(name)
            counts[c.verdict if c else Verdict.NOT_COMPARED] += 1
        summary.append([name, *[counts[v] for v in Verdict]])
    if run_meta:
        summary.append([])
        summary.append(["Run detail", "Value"])
        for key, value in run_meta.items():
            summary.append([key, str(value)])
    _style_workbook_sheet(summary, 8, [])

    workbook.save(path)
    return path


def write_scraped_workbook(listings: Sequence[Any], path: str) -> str:
    """Plain dump of what was scraped, for the scrape-only command."""
    from openpyxl import Workbook

    rows = [l.to_row() for l in listings]
    headers = list(rows[0].keys()) if rows else ["source_url"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Scraped"
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(h) for h in headers])
    _style_workbook_sheet(sheet, len(headers), [])
    workbook.save(path)
    return path


def load_column_map(spec: Optional[str]) -> Optional[Dict[str, int]]:
    """--column-map '{"url": 2, "business_name": 0}' (0-based column indexes)."""
    if not spec:
        return None
    data = json.loads(spec)
    return {str(k): int(v) for k, v in data.items()}
