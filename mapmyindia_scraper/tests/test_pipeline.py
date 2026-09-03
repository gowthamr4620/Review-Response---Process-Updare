"""
End-to-end: spreadsheet in -> scrape -> verify -> report out, with the network
faked at the HTTP boundary so the whole pipeline is exercised offline.
"""

import os
import tempfile
import unittest

from ..excel_io import read_input_rows, write_verification_report
from ..http_client import FetchResult
from ..models import Verdict
from ..scraper import ListingScraper
from ..verify import verify_listing

FIXTURES = os.path.join(os.path.dirname(__file__), "fixtures")


class FakeHttp:
    """Stands in for HttpClient: serves fixtures, records what was requested."""

    def __init__(self, pages):
        self.pages = pages
        self.requested = []

    def get(self, url, **kwargs):
        self.requested.append(url)
        if url not in self.pages:
            return FetchResult(url=url, final_url=url, status=404, text="not found")
        return FetchResult(url=url, final_url=url, status=200, text=self.pages[url], content_type="text/html")

    def allowed(self, url):
        return True


def _fixture(name):
    with open(os.path.join(FIXTURES, name), "r", encoding="utf-8") as f:
        return f.read()


def _write_input_workbook(path):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Store master - Q3", None, None, None, None, None, None])   # banner row
    sheet.append(["Store Name", "Full Address", "PIN Code", "Latitude", "Longitude", "Contact Number", "Mappls Link"])
    sheet.append([
        "Kalyan Jewellers Connaught Place", "Shop No 5, Blk A, Connaught Pl, New Delhi",
        110001, 28.6315, 77.2167, "+91 98765 43210", "https://mappls.com/listing-a",
    ])
    sheet.append([
        "Muthoot Finance Kochi Main", "2nd Flr, Marine Drive, Ernakulam",
        682099, 9.9816, 76.2999, "9812345678", "https://mappls.com/listing-b",
    ])
    workbook.save(path)


class TestPipeline(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.input_path = os.path.join(self.tmp.name, "stores.xlsx")
        self.report_path = os.path.join(self.tmp.name, "report.xlsx")
        _write_input_workbook(self.input_path)
        self.http = FakeHttp(
            {
                "https://mappls.com/listing-a": _fixture("listing_jsonld.html"),
                "https://mappls.com/listing-b": _fixture("listing_nextdata.html"),
            }
        )

    def tearDown(self):
        self.tmp.cleanup()

    def test_columns_are_detected_under_a_banner_row(self):
        rows, mapping, header_index = read_input_rows(self.input_path)
        self.assertEqual(header_index, 1)
        self.assertEqual(mapping["url"], 6)
        self.assertEqual(mapping["business_name"], 0)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0].expected["pincode"], "110001")
        self.assertEqual(rows[0].expected["phone"], "9876543210")

    def test_full_run_produces_a_report_with_the_expected_verdicts(self):
        rows, _mapping, _header = read_input_rows(self.input_path)
        scraper = ListingScraper(http=self.http, use_api=False, verbose=False)

        comparisons = []
        for row in rows:
            listing = scraper.scrape(row.url)
            comparisons.append(verify_listing(row.row_number, row.expected, listing))

        first, second = comparisons
        self.assertEqual(first.row_verdict, "MATCH")
        # The sheet's PIN for the second store is wrong (682099 vs 682031).
        self.assertEqual(second.comparisons["pincode"].verdict, Verdict.MISMATCH)
        self.assertEqual(second.row_verdict, "MISMATCH")

        write_verification_report(comparisons, self.report_path, run_meta={"input file": self.input_path})
        self.assertTrue(os.path.exists(self.report_path))

        from openpyxl import load_workbook

        workbook = load_workbook(self.report_path)
        self.assertEqual(workbook.sheetnames, ["Verification", "Needs Attention", "Summary"])
        sheet = workbook["Verification"]
        headers = [c.value for c in sheet[1]]
        self.assertIn("PIN Code (Sheet)", headers)
        self.assertIn("PIN Code (Mappls)", headers)
        self.assertEqual(sheet.cell(row=2, column=headers.index("Row Verdict") + 1).value, "MATCH")
        self.assertEqual(sheet.cell(row=3, column=headers.index("PIN Code (Mappls)") + 1).value, "682031")
        # Only the failing store is repeated on the triage sheet.
        self.assertEqual(workbook["Needs Attention"].max_row, 2)

    def test_a_dead_link_is_reported_not_crashed(self):
        scraper = ListingScraper(http=self.http, use_api=False, verbose=False)
        listing = scraper.scrape("https://mappls.com/listing-missing")
        self.assertIsNotNone(listing.error)
        self.assertIn("404", listing.error)


if __name__ == "__main__":
    unittest.main()
